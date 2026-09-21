"""
Export and formatting utilities for Syllabus-to-TOC mappings.
Generates styled Excel workbooks and machine-readable JSON files emphasizing
Main Chapters, Chapter Page Ranges, and Content Extraction Page Ranges.
"""

import os
import json
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def sanitize_excel_cell(val):
    """Prevents illegal XML characters in Excel cells."""
    if pd.isna(val) or val is None:
        return ""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub("", str(val))


def export_to_excel(
    df: pd.DataFrame,
    output_path: str,
    title: str = "Syllabus → Textbook TOC Topic & Page Range Mapping",
    model_name: str = "nomic-embed-text",
) -> None:
    """
    Exports TOC alignment results to a styled Excel workbook.
    """
    if df.empty:
        print("Warning: Empty DataFrame provided for Excel export.")
        return

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Syllabus-TOC Mapping"

    # Display columns in user-friendly order and headers
    column_display_names = {
        "syllabus_id": "Topic ID",
        "subject_code": "Subject Code",
        "subject_name": "Subject Name",
        "module_no": "Unit",
        "module_title": "Unit Title",
        "syllabus_topic": "Syllabus Topic",
        "syllabus_subtopic": "Syllabus Subtopic",
        "book_stem": "Textbook",
        "main_chapter": "Main Chapter",
        "chapter_page_range": "Chapter Pages",
        "matched_toc_subtopic": "Matched TOC Subtopic",
        "subtopic_level": "Level",
        "extract_page_range": "Content Extract Pages",
        "similarity_score": "Similarity",
        "confidence": "Confidence",
    }

    cols_to_use = [c for c in column_display_names.keys() if c in df.columns]
    num_cols = len(cols_to_use)

    # 1. Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    subtitle = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Embedding Model: {model_name} | Total Mapped Topics: {len(df)}"
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    sub_cell = ws.cell(row=2, column=1, value=subtitle)
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="595959")
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")

    # 2. Table Headers
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx, col_key in enumerate(cols_to_use, start=1):
        cell = ws.cell(row=4, column=col_idx, value=column_display_names[col_key])
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 3. Confidence Fills
    fill_high = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    font_high = Font(name="Calibri", size=10, bold=True, color="0F5132")

    fill_medium = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    font_medium = Font(name="Calibri", size=10, bold=True, color="664D03")

    fill_low = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid")
    font_low = Font(name="Calibri", size=10, color="41464B")

    fill_below = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    font_below = Font(name="Calibri", size=10, italic=True, color="842029")

    # Chapter / Extraction highlight fills
    fill_extract = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
    font_extract = Font(name="Calibri", size=10, bold=True, color="1A73E8")

    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # 4. Populate Data Rows
    current_row = 5
    for _, row in df.iterrows():
        is_alt = (current_row % 2 == 0)

        for col_idx, col_key in enumerate(cols_to_use, start=1):
            val = row.get(col_key, "")
            cell = ws.cell(row=current_row, column=col_idx, value=sanitize_excel_cell(val))
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            if is_alt:
                cell.fill = alt_fill

            # Alignments
            if col_key in {"syllabus_id", "subject_code", "module_no", "subtopic_level", "chapter_page_range"}:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_key == "extract_page_range":
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill = fill_extract
                cell.font = font_extract
            elif col_key == "similarity_score":
                cell.alignment = Alignment(horizontal="right", vertical="center")
                try:
                    cell.value = float(val)
                    cell.number_format = "0.0000"
                except Exception:
                    pass
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Style confidence column
            if col_key == "confidence":
                conf_val = str(val).strip()
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if conf_val == "High":
                    cell.fill = fill_high
                    cell.font = font_high
                elif conf_val == "Medium":
                    cell.fill = fill_medium
                    cell.font = font_medium
                elif conf_val == "Low":
                    cell.fill = fill_low
                    cell.font = font_low
                elif conf_val == "Below Threshold":
                    cell.fill = fill_below
                    cell.font = font_below

        current_row += 1

    # Auto-fit column widths
    for col_idx, col_key in enumerate(cols_to_use, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(column_display_names[col_key]),
            *(len(str(ws.cell(r, col_idx).value or "")) for r in range(4, min(current_row, 60)))
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

    ws.freeze_panes = "A5"
    wb.save(output_path)
    print(f"Exported styled TOC mapping Excel to: {output_path}")


def export_to_json(df: pd.DataFrame, output_path: str) -> None:
    """Exports hierarchical JSON mapping with TOC chapters and extraction page ranges."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    grouped: Dict[str, Any] = {}

    for _, row in df.iterrows():
        sub_code = row.get("subject_code", "UNKNOWN")
        mod_no = str(row.get("module_no", "0"))
        top_id = row.get("syllabus_id", "")

        if sub_code not in grouped:
            grouped[sub_code] = {
                "subject_code": sub_code,
                "subject_name": row.get("subject_name", ""),
                "modules": {}
            }

        if mod_no not in grouped[sub_code]["modules"]:
            grouped[sub_code]["modules"][mod_no] = {
                "module_no": mod_no,
                "module_title": row.get("module_title", ""),
                "topics": {}
            }

        topics_dict = grouped[sub_code]["modules"][mod_no]["topics"]
        if top_id not in topics_dict:
            topics_dict[top_id] = {
                "id": top_id,
                "topic": row.get("syllabus_topic", ""),
                "sub_topic": row.get("syllabus_subtopic", ""),
                "matches": []
            }

        topics_dict[top_id]["matches"].append({
            "book": row.get("book_stem", ""),
            "main_chapter": row.get("main_chapter", ""),
            "chapter_page_range": row.get("chapter_page_range", ""),
            "matched_toc_subtopic": row.get("matched_toc_subtopic", ""),
            "subtopic_level": row.get("subtopic_level", ""),
            "extract_page_range": row.get("extract_page_range", ""),
            "extract_start_page": row.get("extract_start_page", ""),
            "extract_end_page": row.get("extract_end_page", ""),
            "similarity": row.get("similarity_score", 0.0),
            "confidence": row.get("confidence", ""),
        })

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(grouped, f, indent=2, ensure_ascii=False)
    print(f"Exported JSON TOC mapping to: {output_path}")
